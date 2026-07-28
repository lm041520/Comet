"""RAG 结构化解析基础回归测试。"""

import io
from pathlib import Path

import fitz
import pytest
from docling.datamodel.base_models import InputFormat
from docx import Document as DocxDocument

from app.config import settings
from app.core.exceptions import BizError
from app.core.rag.chunking import chunk_blocks, chunk_parent_child
from app.core.rag.ingestion import (
    BlockType,
    DocumentBlock,
    ParseStatus,
    parse_document_structured,
)
from app.core.rag.ingestion.models import DocumentProfile
from app.core.rag.ingestion.parsers.base import BaseDocumentParser, DocumentParseError
from app.core.rag.ingestion.parsers.docling_parser import DoclingParser, _build_converter
from app.core.rag.ingestion.router import ParserRouter
from app.core.rag.indexing import build_chunk_doc


@pytest.mark.parametrize(
    ("file_ext", "content", "expected"),
    [
        (".txt", "彗记文本".encode(), "彗记文本"),
        (".md", "# 标题\n\n正文".encode(), "标题"),
        (".html", b"<h1>Title</h1><script>bad()</script><p>Body</p>", "Body"),
    ],
)
def test_plain_parsers_expose_text_view(
    file_ext: str,
    content: bytes,
    expected: str,
) -> None:
    parsed = parse_document_structured(file_ext, content)

    assert parsed.status == ParseStatus.OK
    assert parsed.blocks
    assert expected in parsed.text


def test_docx_parser() -> None:
    document = DocxDocument()
    document.add_heading("结构化标题", level=1)
    document.add_paragraph("结构化正文")
    stream = io.BytesIO()
    document.save(stream)

    parsed = parse_document_structured(".docx", stream.getvalue())

    assert parsed.parser_name == "plain"
    assert "结构化标题" in parsed.text
    assert "结构化正文" in parsed.text
    assert parsed.blocks[0].block_type == BlockType.HEADING
    assert parsed.blocks[1].heading_path == ["结构化标题"]


def test_pymupdf_parser_keeps_page_metadata() -> None:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Comet PDF")
    content = document.tobytes()
    document.close()

    parsed = parse_document_structured(".pdf", content)

    assert parsed.parser_name == "pymupdf"
    assert parsed.profile is not None
    assert parsed.profile.scanned_pdf_risk is False
    assert parsed.profile.features["advanced_parser_recommended"] is False
    assert parsed.blocks[0].block_type == BlockType.PARAGRAPH
    assert parsed.blocks[0].page_start == 1
    assert "Comet PDF" in parsed.text


def test_empty_content_has_explicit_error() -> None:
    with pytest.raises(BizError, match="解析结果为空"):
        parse_document_structured(".txt", b"")


def test_unsupported_extension_has_explicit_error() -> None:
    with pytest.raises(BizError, match="不支持的文件类型"):
        parse_document_structured(".xlsx", b"data")


def test_chunk_parent_child_keeps_sentence_content() -> None:
    text = "第一句话。第二句话。"

    current = chunk_parent_child(text)
    assert current
    assert "第一句话" in current[0].content
    assert current[0].children


def test_block_chunker_keeps_trace_metadata_and_protects_table() -> None:
    heading = DocumentBlock(
        content="费用说明",
        block_type=BlockType.HEADING,
        page_start=1,
        page_end=1,
        heading_path=["费用说明"],
    )
    paragraph = DocumentBlock(
        content="这是费用说明正文。",
        page_start=1,
        page_end=1,
        heading_path=["费用说明"],
    )
    table = DocumentBlock(
        content="名称 | 金额\n服务费 | 100",
        block_type=BlockType.TABLE,
        page_start=2,
        page_end=2,
        heading_path=["费用说明"],
    )

    parents = chunk_blocks([heading, paragraph, table])

    assert len(parents) == 2
    assert parents[0].block_ids == [heading.block_id, paragraph.block_id]
    assert parents[0].heading_path == ["费用说明"]
    assert parents[1].block_types == ["table"]
    assert parents[1].page_start == 2
    assert parents[1].child_chunks[0].block_ids == [table.block_id]


def test_chunk_document_contains_block_trace_fields() -> None:
    chunk = build_chunk_doc(
        user_id="user",
        kb_id="kb",
        source_type="document",
        source_id="doc",
        doc_name="demo.pdf",
        chunk_type="child",
        content="命中文本",
        vector=[0.1],
        block_ids=["block-1"],
        block_types=["paragraph"],
        page_start=3,
        page_end=3,
        heading_path=["第三章"],
        chunk_index=2,
        parser_name="pymupdf",
    )["_source"]

    assert chunk["block_ids"] == ["block-1"]
    assert chunk["page_start"] == 3
    assert chunk["heading_path"] == ["第三章"]
    assert chunk["chunk_index"] == 2
    assert chunk["parser_name"] == "pymupdf"


def test_router_uses_docling_only_for_complex_pdf(monkeypatch) -> None:
    monkeypatch.setattr(DoclingParser, "is_available", lambda self: True)
    router = ParserRouter()
    ordinary = DocumentProfile(file_ext=".pdf", file_size=10)
    complex_pdf = DocumentProfile(
        file_ext=".pdf",
        file_size=10,
        features={"advanced_parser_recommended": True},
    )

    assert router.select(ordinary).name == "pymupdf"
    assert router.select(complex_pdf).name == "docling"
    assert complex_pdf.features["route_reason"] == "complex_pdf"


def test_docling_converter_uses_local_onnx_artifacts() -> None:
    _build_converter.cache_clear()
    converter = _build_converter()
    options = converter.format_to_options[InputFormat.PDF].pipeline_options

    assert Path(options.artifacts_path) == Path(settings.rag_docling_cache_dir).resolve()
    assert options.ocr_options.backend == "onnxruntime"


def test_docling_failure_falls_back_to_pymupdf() -> None:
    class FailingDoclingParser(BaseDocumentParser):
        name = "docling"
        supported_exts = frozenset({".pdf"})

        def parse(self, content: bytes, profile: DocumentProfile):
            raise DocumentParseError("模拟 Docling 失败")

    class FailingRouter:
        def select(
            self,
            profile: DocumentProfile,
            preferred_parser: str | None = None,
        ):
            return FailingDoclingParser()

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Fallback PDF")
    content = document.tobytes()
    document.close()

    parsed = parse_document_structured(
        ".pdf",
        content,
        router=FailingRouter(),
    )

    assert parsed.status == ParseStatus.FALLBACK
    assert parsed.parser_name == "pymupdf"
    assert parsed.parse_summary["fallback_from"] == "docling"
    assert "Docling 失败" in parsed.warnings[0]
